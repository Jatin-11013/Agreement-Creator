import streamlit as st
from docx import Document
from datetime import date
import io
import os

st.set_page_config(layout="wide")
st.title("Agreement Generator")

# -----------------------------
# INPUTS
# -----------------------------

col1, col2 = st.columns(2)

with col1:
    org_name = st.text_input("Organisation Name")
    doc_date = st.date_input("Date", value=date.today())
    doc_type = st.text_input("Document Type")
    doc_number = st.text_input("Document Number")
    email = st.text_input("Email")

with col2:
    address = st.text_area("Address")
    org_sign_name = st.text_input("Organisation Signing Name")
    aer_sign_name = st.text_input("Aertrip Signing Name")

# -----------------------------
# DESIGNATION DROPDOWN
# -----------------------------

designation_options = [
    "Director",
    "Partner",
    "Managing Partner",
    "Proprietor",
    "Vice President - Operations",
    "Other"
]

org_designation = st.selectbox("Organisation Signing Designation", designation_options)
if org_designation == "Other":
    org_designation = st.text_input("Enter Organisation Designation")

aer_designation = st.selectbox("Aertrip Signing Designation", designation_options)
if aer_designation == "Other":
    aer_designation = st.text_input("Enter Aertrip Designation")

# -----------------------------
# ANNEXURE CONTROLS
# -----------------------------

col3, col4 = st.columns(2)

with col3:
    annexure_a = st.selectbox("Annexure A: Fees", ["Yes", "No"])

with col4:
    annexure_b = st.selectbox("Annexure B: Related Parties", ["Yes", "No"])

# -----------------------------
# ANNEXURE B INPUTS
# -----------------------------

party_names = []

if annexure_b == "Yes":
    num_parties = st.number_input("Total Number of Parties", min_value=1, step=1)

    for i in range(int(num_parties)):
        name = st.text_input(f"Related Party {i+1} Name", key=f"party_{i}")
        party_names.append(name)

# -----------------------------
# GENERATE BUTTON
# -----------------------------

if st.button("Generate Agreement"):

    try:
        BASE_DIR = os.path.dirname(os.path.abspath(__file__))
        template_path = os.path.join(BASE_DIR, "template.docx")
        annex_path = os.path.join(BASE_DIR, "annexure_a.docx")

        doc = Document(template_path)

        # -----------------------------
        # TEXT REPLACEMENT (FORMATTING SAFE)
        # -----------------------------
        replacements = {
            "{{org_name}}": org_name,
            "{{date}}": str(doc_date),
            "{{document_type}}": doc_type,
            "{{document_number}}": doc_number,
            "{{address}}": address,
            "{{email}}": email,
            "{{org_sign_name}}": org_sign_name,
            "{{org_sign_designation}}": org_designation,
            "{{aer_sign_name}}": aer_sign_name,
            "{{aer_sign_designation}}": aer_designation,
        }

        for para in doc.paragraphs:
            for key, value in replacements.items():
                for run in para.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, value if value else "")

        # -----------------------------
        # ANNEXURE A LINE
        # -----------------------------
        for para in doc.paragraphs:
            for run in para.runs:
                if "{{ANNEXURE_A_LINE}}" in run.text:
                    if annexure_a == "Yes":
                        run.text = run.text.replace(
                            "{{ANNEXURE_A_LINE}}",
                            "The implications of Aertrip Fees are outlined in Annexure A"
                        )
                    else:
                        run.text = run.text.replace("{{ANNEXURE_A_LINE}}", "")

        # -----------------------------
        # ANNEXURE B LINE
        # -----------------------------
        for para in doc.paragraphs:
            for run in para.runs:
                if "{{ANNEXURE_B_LINE}}" in run.text:
                    if annexure_b == "Yes":
                        run.text = run.text.replace(
                            "{{ANNEXURE_B_LINE}}",
                            "and its related entities as mentioned in Annexure B"
                        )
                    else:
                        run.text = run.text.replace("{{ANNEXURE_B_LINE}}", "")

        # -----------------------------
        # ANNEXURE A INSERT AT PLACEHOLDER
        # -----------------------------
        if annexure_a == "Yes":
            annex_doc = Document(annex_path)

            for para in doc.paragraphs:
                if "{{ANNEXURE_A_SECTION}}" in para.text:

                    para.text = ""

                    parent = para._element.getparent()
                    index = parent.index(para._element)

                    for element in annex_doc.element.body:
                        index += 1
                        parent.insert(index, element)

                    break
        else:
            for para in doc.paragraphs:
                if "{{ANNEXURE_A_SECTION}}" in para.text:
                    para.text = ""

        # -----------------------------
        # ANNEXURE B INSERT AT PLACEHOLDER
        # -----------------------------
        if annexure_b == "Yes":

            for para in doc.paragraphs:
                if "{{ANNEXURE_B_SECTION}}" in para.text:

                    para.text = ""

                    parent = para._element.getparent()
                    index = parent.index(para._element)

                    from docx.oxml import OxmlElement
                    from docx.text.paragraph import Paragraph

                    # Heading
                    new_p = OxmlElement("w:p")
                    parent.insert(index + 1, new_p)
                    paragraph = Paragraph(new_p, doc)
                    paragraph.add_run("ANNEXURE B - CLIENT’S ENTITIES\n")

                    index += 1

                    # List
                    for i, name in enumerate(party_names, 1):
                        if name:
                            new_p = OxmlElement("w:p")
                            parent.insert(index + 1, new_p)
                            paragraph = Paragraph(new_p, doc)
                            paragraph.add_run(f"{i}. {name}")
                            index += 1

                    break
        else:
            for para in doc.paragraphs:
                if "{{ANNEXURE_B_SECTION}}" in para.text:
                    para.text = ""

        # -----------------------------
        # SAVE FILE
        # -----------------------------
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        file_name = f"{org_name}_Agreement.docx"

        st.download_button(
            label="Download Agreement",
            data=buffer,
            file_name=file_name,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        st.success("✅ Agreement Generated Successfully!")

    except Exception as e:
        st.error(f"❌ Error: {e}")
