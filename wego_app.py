!pip install streamlit python-docx
import streamlit as st
from docx import Document
from datetime import date
import io

st.set_page_config(layout="wide")
st.title("Agreement Generator")

# -----------------------------
# INPUTS
# -----------------------------

col1, col2 = st.columns(2)

with col1:
    org_name = st.text_input("Organisation Name")
    doc_date = st.date_input("Date", value=date.today())
    doc_type = st.text_input("Document Type")
    doc_number = st.text_input("Document Number")
    email = st.text_input("Email")

with col2:
    address = st.text_area("Address")
    org_sign_name = st.text_input("Organisation Signing Name")
    aer_sign_name = st.text_input("Aertrip Signing Name")

# -----------------------------
# DESIGNATION DROPDOWN
# -----------------------------

designation_options = [
    "Director",
    "Partner",
    "Managing Partner",
    "Proprietor",
    "Vice President - Operations",
    "Other"
]

org_designation = st.selectbox("Organisation Signing Designation", designation_options)

if org_designation == "Other":
    org_designation = st.text_input("Enter Organisation Designation")

aer_designation = st.selectbox("Aertrip Signing Designation", designation_options)

if aer_designation == "Other":
    aer_designation = st.text_input("Enter Aertrip Designation")

# -----------------------------
# ANNEXURE CONTROLS
# -----------------------------

col3, col4 = st.columns(2)

with col3:
    annexure_a = st.selectbox("Annexure A: Fees", ["Yes", "No"])

with col4:
    annexure_b = st.selectbox("Annexure B: Related Parties", ["Yes", "No"])

# -----------------------------
# ANNEXURE B DYNAMIC INPUTS
# -----------------------------

party_names = []

if annexure_b == "Yes":
    num_parties = st.number_input("Total Number of Parties", min_value=1, step=1)

    for i in range(int(num_parties)):
        name = st.text_input(f"Related Party {i+1} Name", key=f"party_{i}")
        party_names.append(name)

# -----------------------------
# GENERATE BUTTON
# -----------------------------

if st.button("Generate Agreement"):

    # Load template
    doc = Document("template.docx")

    # -----------------------------
    # BASIC REPLACEMENTS
    # -----------------------------
    replacements = {
        "{{org_name}}": org_name,
        "{{date}}": str(doc_date),
        "{{document_type}}": doc_type,
        "{{document_number}}": doc_number,
        "{{address}}": address,
        "{{email}}": email,
        "{{org_sign_name}}": org_sign_name,
        "{{org_sign_designation}}": org_designation,
        "{{aer_sign_name}}": aer_sign_name,
        "{{aer_sign_designation}}": aer_designation,
    }

    # Replace text
    for para in doc.paragraphs:
        for key, value in replacements.items():
            if key in para.text:
                para.text = para.text.replace(key, value if value else "")

    # -----------------------------
    # ANNEXURE A LINE
    # -----------------------------
    if annexure_a == "Yes":
        annexure_a_line = "The implications of Aertrip Fees are outlined in Annexure A"
    else:
        annexure_a_line = ""

    for para in doc.paragraphs:
        if "{{ANNEXURE_A_LINE}}" in para.text:
            para.text = para.text.replace("{{ANNEXURE_A_LINE}}", annexure_a_line)

    # -----------------------------
    # ANNEXURE B LINE
    # -----------------------------
    if annexure_b == "Yes":
        annexure_b_line = "and its related entities as mentioned in Annexure B"
    else:
        annexure_b_line = ""

    for para in doc.paragraphs:
        if "{{ANNEXURE_B_LINE}}" in para.text:
            para.text = para.text.replace("{{ANNEXURE_B_LINE}}", annexure_b_line)

    # -----------------------------
    # REMOVE PLACEHOLDER TEXTS
    # -----------------------------
    for para in doc.paragraphs:
        para.text = para.text.replace("{{ANNEXURE_A_SECTION}}", "")
        para.text = para.text.replace("{{ANNEXURE_B_SECTION}}", "")

    # -----------------------------
    # ANNEXURE A INSERT
    # -----------------------------
    if annexure_a == "Yes":
        annex_doc = Document("annexure_a.docx")

        doc.add_page_break()

        for element in annex_doc.element.body:
            doc.element.body.append(element)

    # -----------------------------
    # ANNEXURE B INSERT
    # -----------------------------
    if annexure_b == "Yes":
        doc.add_page_break()
        doc.add_paragraph("ANNEXURE B - CLIENT’S ENTITIES\n")

        for i, name in enumerate(party_names, 1):
            if name:
                doc.add_paragraph(f"{i}. {name}")

    # -----------------------------
    # SAVE FILE (IN MEMORY)
    # -----------------------------
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # -----------------------------
    # DOWNLOAD BUTTON
    # -----------------------------
    file_name = f"{org_name}_Agreement.docx"

    st.download_button(
        label="Download Agreement",
        data=buffer,
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
